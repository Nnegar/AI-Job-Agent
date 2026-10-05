import sys
from pathlib import Path
import docx

RESUMES_DIR = Path(__file__).resolve().parents[1] / "resumes"

def docx_to_markdown(docx_path: Path) -> str:
    doc = docx.Document(docx_path)
    md_lines = []
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
            
        style_name = p.style.name.lower() if p.style else ""
        
        if "heading 1" in style_name or "title" in style_name:
            md_lines.append(f"# {text}\n")
        elif "heading 2" in style_name:
            md_lines.append(f"## {text}\n")
        elif "heading 3" in style_name:
            md_lines.append(f"### {text}\n")
        elif "bullet" in style_name or "list" in style_name:
            cleaned = text.lstrip("•- \t").strip()
            md_lines.append(f"- {cleaned}")
        else:
            # Check if paragraph looks like a list item or section header
            if text.startswith("•") or text.startswith("-"):
                cleaned = text.lstrip("•- \t").strip()
                md_lines.append(f"- {cleaned}")
            elif text.isupper() and len(text) < 40:
                md_lines.append(f"\n## {text.title()}\n")
            else:
                md_lines.append(f"{text}\n")

    # Also extract tables if any
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(row_cells):
                table_rows.append(row_cells)
        if table_rows:
            md_lines.append("\n")
            for r in table_rows:
                md_lines.append("| " + " | ".join(r) + " |")
            md_lines.append("\n")
            
    return "\n".join(md_lines).strip()

def main():
    docx_files = sorted(RESUMES_DIR.glob("*.docx"))
    if not docx_files:
        print(f"No .docx files found in {RESUMES_DIR}")
        return

    for docx_path in docx_files:
        if ":Zone.Identifier" in docx_path.name:
            continue
        md_name = docx_path.stem + ".md"
        md_path = RESUMES_DIR / md_name
        text = docx_to_markdown(docx_path)
        md_path.write_text(text, encoding="utf-8")
        print(f"Extracted {docx_path.name} -> {md_name} ({len(text)} chars, {len(text.splitlines())} lines)")

if __name__ == "__main__":
    main()
